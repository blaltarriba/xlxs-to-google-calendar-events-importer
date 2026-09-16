"""Tests for reading one group's training details out of the trimester sheet."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from golf_calendar.domain import TrainingDetail
from golf_calendar.training_details_reader import TrainingDetailsReadError
from golf_calendar.training_details_reader import TrainingDetailsValidationError
from golf_calendar.training_details_reader import read_training_details

REAL_DETAILS = Path(__file__).parents[2] / "Entrenamientos martes 1 trim.xlsx"

GROUPS = ("Grupo A", "Grupo E")


class DetailsSheetBuilder:
    """Builds a trimester sheet shaped like the school's, stating only what a test cares about."""

    def __init__(self, header_row: int = 4) -> None:
        self.workbook = Workbook()
        self.sheet: Worksheet = self.workbook.worksheets[0]
        self.sheet.cell(row=1, column=1, value="MARTES · ENTRENAMIENTOS")
        self._next_row = header_row + 1
        for column, header in enumerate(("Fecha", *GROUPS, "Observación"), start=1):
            self.sheet.cell(row=header_row, column=column, value=header)

    def with_row(
        self, day: date | str | None, group_a: str | None, group_e: str | None
    ) -> DetailsSheetBuilder:
        self.sheet.cell(row=self._next_row, column=1, value=day)
        self.sheet.cell(row=self._next_row, column=2, value=group_a)
        self.sheet.cell(row=self._next_row, column=3, value=group_e)
        self._next_row += 1
        return self

    def save(self, tmp_path: Path) -> Path:
        path = tmp_path / "details.xlsx"
        self.workbook.save(path)
        return path


def detail(day: date, activities: str) -> TrainingDetail:
    return TrainingDetail(session_date=day, group_label="Grupo E", activities=activities)


class TestReadsTheRealSpreadsheet:
    def test_reads_every_group_e_detail_of_the_first_trimester(self) -> None:
        details = read_training_details(REAL_DETAILS, "Grupo E")

        assert details == (
            detail(date(2026, 9, 15), "Putt + Juego largo"),
            detail(date(2026, 9, 22), "Approach + Putt"),
            detail(date(2026, 9, 29), "P&P estrategia + Juego largo"),
            detail(date(2026, 10, 6), "Approach + Juego largo"),
            detail(date(2026, 10, 20), "Putt + Approach"),
            detail(date(2026, 10, 27), "P&P estrategia + Putt"),
            detail(date(2026, 11, 3), "Approach + Putt"),
            detail(date(2026, 11, 10), "Juego largo + Putt"),
            detail(date(2026, 11, 17), "Approach + Juego largo"),
            detail(date(2026, 11, 24), "P&P estrategia + Approach"),
            detail(date(2026, 12, 1), "Juego largo + Approach"),
            detail(date(2026, 12, 15), "P&P estrategia + Approach"),
        )


class TestReadsABuiltSpreadsheet:
    def test_reads_only_the_chosen_groups_column(self, tmp_path: Path) -> None:
        details_file = (
            DetailsSheetBuilder()
            .with_row(date(2026, 9, 15), "Approach", "  Putt + Juego largo ")
            .save(tmp_path)
        )

        details = read_training_details(details_file, "Grupo E")

        assert details == (detail(date(2026, 9, 15), "Putt + Juego largo"),)

    @pytest.mark.parametrize("configured_group", ["grupo e", "  GRUPO E  ", "Grupo  E"])
    def test_matches_the_group_ignoring_case_and_spacing(
        self, tmp_path: Path, configured_group: str
    ) -> None:
        details_file = (
            DetailsSheetBuilder().with_row(date(2026, 9, 15), "Approach", "Putt").save(tmp_path)
        )

        details = read_training_details(details_file, configured_group)

        assert details == (detail(date(2026, 9, 15), "Putt"),)

    def test_skips_rows_without_a_date_and_orders_by_date(self, tmp_path: Path) -> None:
        details_file = (
            DetailsSheetBuilder()
            .with_row(date(2026, 9, 22), "Approach", "Approach + Putt")
            .with_row(None, None, None)
            .with_row(date(2026, 9, 15), "Approach", "Putt + Juego largo")
            .with_row(None, None, None)
            .save(tmp_path)
        )

        details = read_training_details(details_file, "Grupo E")

        assert details == (
            detail(date(2026, 9, 15), "Putt + Juego largo"),
            detail(date(2026, 9, 22), "Approach + Putt"),
        )

    def test_finds_the_header_on_a_later_sheet_and_row(self, tmp_path: Path) -> None:
        builder = DetailsSheetBuilder(header_row=9).with_row(date(2026, 9, 15), "A", "Putt")
        cover = builder.workbook.create_sheet("Portada", index=0)
        cover.cell(row=1, column=1, value="Escuela de Golf")

        details = read_training_details(builder.save(tmp_path), "Grupo E")

        assert details == (detail(date(2026, 9, 15), "Putt"),)

    def test_reads_an_empty_trimester_as_no_details(self, tmp_path: Path) -> None:
        details_file = DetailsSheetBuilder().save(tmp_path)

        assert read_training_details(details_file, "Grupo E") == ()


class TestRejectsAnUnusableSpreadsheet:
    def test_a_file_that_cannot_be_opened(self, tmp_path: Path) -> None:
        missing = tmp_path / "missing.xlsx"

        with pytest.raises(TrainingDetailsReadError, match="cannot open training details"):
            read_training_details(missing, "Grupo E")

    def test_a_sheet_without_a_date_header(self, tmp_path: Path) -> None:
        workbook = Workbook()
        workbook.worksheets[0].cell(row=1, column=1, value="Grupo E")
        details_file = tmp_path / "details.xlsx"
        workbook.save(details_file)

        with pytest.raises(TrainingDetailsReadError, match="no 'Fecha' header"):
            read_training_details(details_file, "Grupo E")

    def test_names_the_groups_available_when_the_group_is_missing(self, tmp_path: Path) -> None:
        details_file = DetailsSheetBuilder().save(tmp_path)

        with pytest.raises(TrainingDetailsReadError, match="available: Grupo A, Grupo E"):
            read_training_details(details_file, "Grupo F")

    def test_a_date_column_holding_text(self, tmp_path: Path) -> None:
        details_file = DetailsSheetBuilder().with_row("15/09", "A", "Putt").save(tmp_path)

        with pytest.raises(TrainingDetailsValidationError, match="'15/09' is not a date"):
            read_training_details(details_file, "Grupo E")

    def test_a_dated_row_without_a_detail_for_the_group(self, tmp_path: Path) -> None:
        details_file = DetailsSheetBuilder().with_row(date(2026, 9, 15), "A", "   ").save(tmp_path)

        with pytest.raises(TrainingDetailsValidationError, match="no Grupo E detail on 2026-09-15"):
            read_training_details(details_file, "Grupo E")

    def test_a_date_listed_twice(self, tmp_path: Path) -> None:
        details_file = (
            DetailsSheetBuilder()
            .with_row(date(2026, 9, 15), "A", "Putt")
            .with_row(date(2026, 9, 15), "A", "Approach")
            .save(tmp_path)
        )

        with pytest.raises(TrainingDetailsValidationError, match="2026-09-15 is listed twice"):
            read_training_details(details_file, "Grupo E")
