"""Tests for reading the season out of the colour-coded spreadsheet."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from golf_calendar.domain import TrainingSession
from golf_calendar.domain import TrainingWeekday
from golf_calendar.excel_schedule_reader import ScheduleReadError
from golf_calendar.excel_schedule_reader import ScheduleValidationError
from golf_calendar.excel_schedule_reader import read_season

REAL_SCHEDULE = Path(__file__).parents[2] / "Calendario 2627 L,M,M.xlsx"

LEGEND_COLOURS = {
    TrainingWeekday.MONDAY: "5B9BD5",
    TrainingWeekday.TUESDAY: "E06666",
    TrainingWeekday.WEDNESDAY: "70AD47",
}
HOLIDAY_COLOUR = "D9D9D9"

EXPECTED_TUESDAYS = (
    date(2026, 9, 15),
    date(2026, 9, 22),
    date(2026, 9, 29),
    date(2026, 10, 6),
    date(2026, 10, 20),
    date(2026, 10, 27),
    date(2026, 11, 3),
    date(2026, 11, 10),
    date(2026, 11, 17),
    date(2026, 11, 24),
    date(2026, 12, 1),
    date(2026, 12, 15),
    date(2027, 1, 12),
    date(2027, 1, 19),
    date(2027, 1, 26),
    date(2027, 2, 9),
    date(2027, 2, 16),
    date(2027, 2, 23),
    date(2027, 3, 2),
    date(2027, 3, 9),
    date(2027, 3, 23),
    date(2027, 4, 13),
    date(2027, 4, 20),
    date(2027, 4, 27),
    date(2027, 5, 4),
    date(2027, 5, 11),
    date(2027, 5, 18),
    date(2027, 5, 25),
    date(2027, 6, 1),
    date(2027, 6, 8),
)


def fill(colour: str) -> PatternFill:
    return PatternFill(patternType="solid", start_color=f"FF{colour}", end_color=f"FF{colour}")


class ScheduleBuilder:
    """Builds a spreadsheet shaped like the school's, stating only what a test cares about."""

    def __init__(self) -> None:
        self.workbook = Workbook()
        self.sheet: Worksheet = self.workbook.worksheets[0]
        self._next_row = 1
        self._last_heading_row = 1
        self._write_legend()

    def _write_legend(self) -> None:
        self.sheet.cell(row=1, column=1, value="LEYENDA")
        for index, (weekday, colour) in enumerate(LEGEND_COLOURS.items()):
            cell = self.sheet.cell(row=2, column=1 + index, value=weekday.legend_label)
            cell.fill = fill(colour)
        self._next_row = 4

    def with_subtitle(self, text: str) -> ScheduleBuilder:
        self.sheet.cell(row=3, column=1, value=text)
        return self

    def with_month(
        self,
        year: int,
        heading: str,
        weeks: list[list[int | None]],
        sessions: dict[int, str],
        first_column: int = 1,
        weekday_initials: str = "LMXJVSD",
    ) -> ScheduleBuilder:
        """Add a month grid below the last. ``sessions`` maps a day number to its fill colour."""
        return self._write_month(
            self._next_row, year, heading, weeks, sessions, first_column, weekday_initials
        )

    def with_month_beside(
        self,
        year: int,
        heading: str,
        weeks: list[list[int | None]],
        sessions: dict[int, str],
        first_column: int,
    ) -> ScheduleBuilder:
        """Add a month grid alongside the previous one, as the school's sheet lays them out."""
        return self._write_month(
            self._last_heading_row, year, heading, weeks, sessions, first_column, "LMXJVSD"
        )

    def _write_month(
        self,
        heading_row: int,
        year: int,
        heading: str,
        weeks: list[list[int | None]],
        sessions: dict[int, str],
        first_column: int,
        weekday_initials: str,
    ) -> ScheduleBuilder:
        self._last_heading_row = heading_row
        self.sheet.cell(row=heading_row, column=first_column, value=f"{heading} {year}")
        for offset, initial in enumerate(weekday_initials):
            self.sheet.cell(row=heading_row + 1, column=first_column + offset, value=initial)
        for week_index, week in enumerate(weeks):
            row = heading_row + 2 + week_index
            for offset, day in enumerate(week):
                if day is None:
                    continue
                cell = self.sheet.cell(row=row, column=first_column + offset, value=day)
                if day in sessions:
                    cell.fill = fill(sessions[day])
        self._next_row = max(self._next_row, heading_row + 2 + len(weeks) + 1)
        return self

    def save(self, tmp_path: Path) -> Path:
        path = tmp_path / "schedule.xlsx"
        self.workbook.save(path)
        return path


def september_2026() -> list[list[int | None]]:
    """September 2026 laid out Monday-first, exactly as the school's sheet lays it out."""
    return [
        [None, 1, 2, 3, 4, 5, 6],
        [7, 8, 9, 10, 11, 12, 13],
        [14, 15, 16, 17, 18, 19, 20],
        [21, 22, 23, 24, 25, 26, 27],
        [28, 29, 30, None, None, None, None],
    ]


class TestReadsTheRealSpreadsheet:
    def test_cross_checks_against_the_sheets_own_declared_total(self) -> None:
        season = read_season(
            REAL_SCHEDULE, TrainingWeekday.TUESDAY, extra_session_dates=(date(2027, 6, 1),)
        )

        assert season.declared_total == 30
        assert season.key == "2026-2027"

    def test_reads_the_whole_tuesday_season(self) -> None:
        season = read_season(
            REAL_SCHEDULE, TrainingWeekday.TUESDAY, extra_session_dates=(date(2027, 6, 1),)
        )

        assert tuple(s.session_date for s in season.sessions) == EXPECTED_TUESDAYS
        assert season.sessions[0] == TrainingSession(date(2026, 9, 15), number=1, total=30)
        assert season.sessions[-1] == TrainingSession(date(2027, 6, 8), number=30, total=30)

    @pytest.mark.parametrize(
        ("weekday", "extra", "first", "last"),
        [
            (TrainingWeekday.MONDAY, (), date(2026, 9, 14), date(2027, 6, 7)),
            (TrainingWeekday.TUESDAY, (date(2027, 6, 1),), date(2026, 9, 15), date(2027, 6, 8)),
            (TrainingWeekday.WEDNESDAY, (date(2027, 6, 2),), date(2026, 9, 16), date(2027, 6, 9)),
        ],
    )
    def test_every_weekday_of_the_season_has_thirty_sessions(
        self,
        weekday: TrainingWeekday,
        extra: tuple[date, ...],
        first: date,
        last: date,
    ) -> None:
        season = read_season(REAL_SCHEDULE, weekday, extra_session_dates=extra)

        assert len(season) == 30
        assert season.sessions[0].session_date == first
        assert season.sessions[-1].session_date == last

    def test_every_session_falls_on_the_requested_weekday(self) -> None:
        season = read_season(
            REAL_SCHEDULE, TrainingWeekday.TUESDAY, extra_session_dates=(date(2027, 6, 1),)
        )

        assert all(TrainingWeekday.TUESDAY.matches(s.session_date) for s in season.sessions)

    def test_excludes_holidays_and_the_recovery_week(self) -> None:
        season = read_season(
            REAL_SCHEDULE, TrainingWeekday.TUESDAY, extra_session_dates=(date(2027, 6, 1),)
        )
        found = {session.session_date for session in season.sessions}

        assert date(2026, 12, 8) not in found, "8 December is a public holiday"
        assert date(2026, 12, 29) not in found, "Navidad"
        assert date(2027, 2, 2) not in found, "margen / recovery week"
        assert date(2027, 3, 16) not in found, "Fallas"
        assert date(2027, 3, 30) not in found, "Pascua"


class TestReadsABuiltSpreadsheet:
    @pytest.mark.parametrize("weekday", list(TrainingWeekday))
    def test_finds_the_days_carrying_that_weekdays_legend_colour(
        self, tmp_path: Path, weekday: TrainingWeekday
    ) -> None:
        colour = LEGEND_COLOURS[weekday]
        sessions = {14: LEGEND_COLOURS[TrainingWeekday.MONDAY]} | {
            15: LEGEND_COLOURS[TrainingWeekday.TUESDAY],
            16: LEGEND_COLOURS[TrainingWeekday.WEDNESDAY],
        }
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), sessions)
            .save(tmp_path)
        )

        season = read_season(schedule, weekday)

        expected_day = {"5B9BD5": 14, "E06666": 15, "70AD47": 16}[colour]
        assert season.sessions == (TrainingSession(date(2026, 9, expected_day), number=1, total=1),)

    def test_reads_nothing_when_no_day_carries_the_colour(self, tmp_path: Path) -> None:
        schedule = (
            ScheduleBuilder().with_month(2026, "SEPTIEMBRE", september_2026(), {}).save(tmp_path)
        )

        assert read_season(schedule, TrainingWeekday.TUESDAY).sessions == ()

    def test_ignores_holiday_and_other_weekday_colours(self, tmp_path: Path) -> None:
        sessions = {
            8: HOLIDAY_COLOUR,
            14: LEGEND_COLOURS[TrainingWeekday.MONDAY],
            15: LEGEND_COLOURS[TrainingWeekday.TUESDAY],
            16: LEGEND_COLOURS[TrainingWeekday.WEDNESDAY],
            22: HOLIDAY_COLOUR,
        }
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), sessions)
            .save(tmp_path)
        )

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert season.sessions == (TrainingSession(date(2026, 9, 15), number=1, total=1),)

    def test_numbers_sessions_across_several_month_grids(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        october = {6: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .with_month(
                2026,
                "OCTUBRE",
                [[None, None, None, 1, 2, 3, 4], [5, 6, 7, 8, 9, 10, 11]],
                october,
            )
            .save(tmp_path)
        )

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert season.sessions == (
            TrainingSession(date(2026, 9, 15), number=1, total=2),
            TrainingSession(date(2026, 10, 6), number=2, total=2),
        )

    def test_reads_month_grids_placed_side_by_side(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        october = {6: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .with_month_beside(
                2026,
                "OCTUBRE",
                [[None, None, None, 1, 2, 3, 4], [5, 6, 7, 8, 9, 10, 11]],
                october,
                first_column=9,
            )
            .save(tmp_path)
        )

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert tuple(s.session_date for s in season.sessions) == (
            date(2026, 9, 15),
            date(2026, 10, 6),
        )


class TestExtraSessionDates:
    def test_merges_a_date_the_sheet_failed_to_colour(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        season = read_season(
            schedule, TrainingWeekday.TUESDAY, extra_session_dates=(date(2026, 9, 22),)
        )

        assert season.sessions == (
            TrainingSession(date(2026, 9, 15), number=1, total=2),
            TrainingSession(date(2026, 9, 22), number=2, total=2),
        )

    def test_rejects_an_extra_date_on_the_wrong_weekday(self, tmp_path: Path) -> None:
        schedule = (
            ScheduleBuilder().with_month(2026, "SEPTIEMBRE", september_2026(), {}).save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="is not a Martes"):
            read_season(schedule, TrainingWeekday.TUESDAY, extra_session_dates=(date(2026, 9, 16),))

    def test_rejects_the_same_extra_date_supplied_twice(self, tmp_path: Path) -> None:
        """Repeating an override must fail with a reason, not a bare numbering crash."""
        schedule = (
            ScheduleBuilder().with_month(2026, "SEPTIEMBRE", september_2026(), {}).save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="already a session"):
            read_season(
                schedule,
                TrainingWeekday.TUESDAY,
                extra_session_dates=(date(2026, 9, 22), date(2026, 9, 22)),
            )

    def test_rejects_an_extra_date_outside_the_months_the_sheet_covers(
        self, tmp_path: Path
    ) -> None:
        """A year typo that still lands on a Tuesday would renumber the whole season."""
        schedule = (
            ScheduleBuilder().with_month(2026, "SEPTIEMBRE", september_2026(), {}).save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="outside the months the sheet covers"):
            read_season(schedule, TrainingWeekday.TUESDAY, extra_session_dates=(date(2025, 9, 16),))

    def test_rejects_an_extra_date_the_sheet_already_colours(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="already a session"):
            read_season(schedule, TrainingWeekday.TUESDAY, extra_session_dates=(date(2026, 9, 15),))


class TestRejectsAnUnreadableSheet:
    def test_missing_legend_colour(self, tmp_path: Path) -> None:
        workbook = Workbook()
        sheet: Worksheet = workbook.worksheets[0]
        sheet.cell(row=1, column=1, value="SEPTIEMBRE 2026")
        path = tmp_path / "schedule.xlsx"
        workbook.save(path)

        with pytest.raises(ScheduleReadError, match="no filled legend cell labelled 'Martes'"):
            read_season(path, TrainingWeekday.TUESDAY)

    def test_no_month_headings(self, tmp_path: Path) -> None:
        path = ScheduleBuilder().save(tmp_path)

        with pytest.raises(ScheduleReadError, match="no month headings found"):
            read_season(path, TrainingWeekday.TUESDAY)

    def test_incomplete_weekday_header_row(self, tmp_path: Path) -> None:
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), {}, weekday_initials="LMXJV??")
            .save(tmp_path)
        )

        with pytest.raises(ScheduleReadError, match="has 5 weekday columns, expected 7"):
            read_season(schedule, TrainingWeekday.TUESDAY)

    def test_a_day_number_that_is_not_a_date_in_that_month(self, tmp_path: Path) -> None:
        sessions = {31: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", [[None, 31, None, None, None, None, None]], sessions)
            .save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="not a date in 2026-09"):
            read_season(schedule, TrainingWeekday.TUESDAY)

    def test_a_day_sitting_under_the_wrong_weekday_column(self, tmp_path: Path) -> None:
        """A shifted grid must fail loudly rather than yield a plausible wrong date."""
        sessions = {14: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", [[None, 14, None, None, None, None, None]], sessions)
            .save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="the grid was misread"):
            read_season(schedule, TrainingWeekday.TUESDAY)


class TestGridBounds:
    def test_ignores_numbers_below_the_last_month_grid(self, tmp_path: Path) -> None:
        """A totals or footnote row under the final grid must not read as more days."""
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        builder = ScheduleBuilder().with_month(2026, "SEPTIEMBRE", september_2026(), september)
        footer_row = builder.sheet.max_row + 2
        builder.sheet.cell(row=footer_row, column=1, value="Total sesiones")
        stray = builder.sheet.cell(row=footer_row, column=2, value=30)
        stray.fill = fill(LEGEND_COLOURS[TrainingWeekday.TUESDAY])

        season = read_season(builder.save(tmp_path), TrainingWeekday.TUESDAY)

        assert season.sessions == (TrainingSession(date(2026, 9, 15), number=1, total=1),)

    def test_reads_a_month_spanning_six_week_rows(self, tmp_path: Path) -> None:
        """August 2026 needs all six rows, so the cap must not truncate a real month."""
        sessions = {4: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        weeks: list[list[int | None]] = [
            [None, None, None, None, None, 1, 2],
            [3, 4, 5, 6, 7, 8, 9],
            [10, 11, 12, 13, 14, 15, 16],
            [17, 18, 19, 20, 21, 22, 23],
            [24, 25, 26, 27, 28, 29, 30],
            [31, None, None, None, None, None, None],
        ]
        schedule = ScheduleBuilder().with_month(2026, "AGOSTO", weeks, sessions).save(tmp_path)

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert season.sessions == (TrainingSession(date(2026, 8, 4), number=1, total=1),)


class TestDeclaredTotal:
    def test_accepts_a_count_matching_the_subtitle(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_subtitle("1 lunes · 1 martes · 1 miércoles")
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        assert len(read_season(schedule, TrainingWeekday.TUESDAY)) == 1

    def test_reports_that_a_sheet_without_a_subtitle_was_not_cross_checked(
        self, tmp_path: Path
    ) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert season.declared_total is None
        assert season.is_cross_checked is False

    def test_reports_the_declared_total_it_checked_against(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_subtitle("1 lunes · 1 martes · 1 miércoles")
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        season = read_season(schedule, TrainingWeekday.TUESDAY)

        assert season.declared_total == 1
        assert season.is_cross_checked is True

    def test_rejects_a_count_the_sheet_contradicts(self, tmp_path: Path) -> None:
        september = {15: LEGEND_COLOURS[TrainingWeekday.TUESDAY]}
        schedule = (
            ScheduleBuilder()
            .with_subtitle("30 lunes · 30 martes · 30 miércoles")
            .with_month(2026, "SEPTIEMBRE", september_2026(), september)
            .save(tmp_path)
        )

        with pytest.raises(ScheduleValidationError, match="declares 30 martes sessions but 1"):
            read_season(schedule, TrainingWeekday.TUESDAY)
