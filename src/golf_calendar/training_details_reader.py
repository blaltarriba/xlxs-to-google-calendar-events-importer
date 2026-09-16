"""Reads what one group practises each session out of the school's trimester sheet.

The sheet is a plain table: a ``Fecha`` column of session dates and one column per group
(``Grupo A`` … ``Grupo E``) saying what that group trains that day. The table is located by
its ``Fecha`` heading rather than by fixed coordinates, so a sheet that gains a title row or
a cover sheet still reads correctly.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from golf_calendar.domain import GolfCalendarError
from golf_calendar.domain import TrainingDetail

DATE_HEADER = "Fecha"

type SheetRow = tuple[object, ...]


class TrainingDetailsReadError(GolfCalendarError):
    """Raised when the sheet cannot be opened or its table or group cannot be located."""

    code = "training_details_unreadable"


class TrainingDetailsValidationError(GolfCalendarError):
    """Raised when the table is found but one of its rows cannot be trusted."""

    code = "training_details_invalid"


def read_training_details(details_file: Path, group_label: str) -> tuple[TrainingDetail, ...]:
    """Return ``group_label``'s detail for every dated row of the sheet, ordered by date.

    The group is matched ignoring case and spacing, so ``grupo e`` in configuration finds
    the sheet's ``Grupo E``. Rows without a date are layout, not sessions, and are skipped.
    """
    header, rows = _find_table(_read_sheets(details_file), details_file)
    date_column = header.index(DATE_HEADER)
    group_column = _find_group_column(header, group_label, details_file)
    details = [
        _read_detail(row, date_column, group_column, header[group_column])
        for row in rows
        if _value_at(row, date_column) is not None
    ]
    _reject_repeated_dates(details)
    return tuple(sorted(details, key=lambda detail: detail.session_date))


def _read_sheets(details_file: Path) -> list[list[SheetRow]]:
    """Load every sheet's cell values at once, so the file is not held open while parsing."""
    try:
        workbook = load_workbook(filename=details_file, read_only=True, data_only=True)
    except (OSError, InvalidFileException, BadZipFile) as error:
        raise TrainingDetailsReadError(
            f"cannot open training details file {details_file}"
        ) from error
    sheets: list[list[SheetRow]] = [
        [tuple(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook.worksheets
    ]
    workbook.close()
    return sheets


def _find_table(
    sheets: list[list[SheetRow]], details_file: Path
) -> tuple[list[str], list[SheetRow]]:
    """The header row as text, and every row below it, from the first sheet that has one."""
    for rows in sheets:
        header_index = _find_header_index(rows)
        if header_index is not None:
            return [_text_of(value) for value in rows[header_index]], rows[header_index + 1 :]
    raise TrainingDetailsReadError(f"no {DATE_HEADER!r} header found in {details_file}")


def _find_header_index(rows: list[SheetRow]) -> int | None:
    for index, row in enumerate(rows):
        if DATE_HEADER in (_text_of(value) for value in row):
            return index
    return None


def _find_group_column(header: list[str], group_label: str, details_file: Path) -> int:
    wanted = _normalised(group_label)
    for index, heading in enumerate(header):
        if heading and _normalised(heading) == wanted:
            return index
    available = ", ".join(heading for heading in header if heading and heading != DATE_HEADER)
    raise TrainingDetailsReadError(
        f"no column {group_label.strip()!r} in {details_file}; available: {available}"
    )


def _read_detail(
    row: SheetRow, date_column: int, group_column: int, sheet_group_label: str
) -> TrainingDetail:
    session_date = _read_date(_value_at(row, date_column))
    activities = _text_of(_value_at(row, group_column))
    if not activities:
        raise TrainingDetailsValidationError(
            f"no {sheet_group_label} detail on {session_date.isoformat()}"
        )
    return TrainingDetail(
        session_date=session_date, group_label=sheet_group_label, activities=activities
    )


def _read_date(value: object) -> date:
    # datetime subclasses date, and a spreadsheet date cell arrives as a datetime.
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise TrainingDetailsValidationError(f"{value!r} is not a date in the {DATE_HEADER!r} column")


def _reject_repeated_dates(details: Sequence[TrainingDetail]) -> None:
    seen: set[date] = set()
    for detail in details:
        if detail.session_date in seen:
            raise TrainingDetailsValidationError(
                f"{detail.session_date.isoformat()} is listed twice"
            )
        seen.add(detail.session_date)


def _value_at(row: SheetRow, column: int) -> object:
    return row[column] if column < len(row) else None


def _text_of(value: object) -> str:
    return value.strip() if isinstance(value, str) else ""


def _normalised(text: str) -> str:
    return " ".join(text.split()).casefold()
