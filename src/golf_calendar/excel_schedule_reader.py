"""Reads the season's training dates out of the school's spreadsheet.

The sheet is a *visual* calendar: a grid of month blocks in which a training day is marked
by the cell's **fill colour**, with a legend row naming each colour. Nothing states a date
in text, so this module reconstructs the dates from the grid's geometry.

It parses structurally rather than by hard-coded coordinates — month blocks are located by
their headings and weekday columns by the ``L M X J V S D`` row beneath them — so a sheet
that gains a month, shifts down a row or is recoloured still reads correctly. Every date it
derives is cross-checked against the column it came from, because a misparsed grid must
fail loudly rather than produce plausible wrong dates.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from golf_calendar.domain import GolfCalendarError
from golf_calendar.domain import Season
from golf_calendar.domain import TrainingWeekday
from golf_calendar.domain import number_sessions

SPANISH_MONTHS = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}

# The initials the sheet uses for its weekday header row, in Spanish: lunes, martes,
# miércoles (X), jueves, viernes, sábado, domingo.
SPANISH_WEEKDAY_INITIALS = {"L": 0, "M": 1, "X": 2, "J": 3, "V": 4, "S": 5, "D": 6}

DAYS_PER_WEEK = 7

# No calendar month spans more than six week rows, so the day grid below a heading is
# bounded even when the heading is the last one on the sheet. Without this, stray numbers
# in a footer row get read as days of the final month.
MAX_WEEK_ROWS = 6

# openpyxl hands back either kind when walking a sheet; both carry a value and a fill.
type AnyCell = Cell | MergedCell

_MONTH_HEADER = re.compile(
    rf"^(?P<month>{'|'.join(SPANISH_MONTHS)})\s+(?P<year>\d{{4}})$",
    re.IGNORECASE,
)


class ScheduleReadError(GolfCalendarError):
    """Raised when the spreadsheet cannot be opened or its grid cannot be located."""

    code = "schedule_unreadable"


class ScheduleValidationError(GolfCalendarError):
    """Raised when the grid parses but says something self-contradictory."""

    code = "schedule_invalid"


@dataclass(frozen=True, slots=True)
class MonthBlock:
    """One month's grid within the sheet: its heading, and the rows of days below it."""

    month: int
    year: int
    heading_row: int
    first_column: int
    first_day_row: int
    last_day_row: int

    @property
    def weekday_header_row(self) -> int:
        return self.heading_row + 1

    @property
    def columns(self) -> range:
        return range(self.first_column, self.first_column + DAYS_PER_WEEK)


def read_season(
    schedule_file: Path,
    weekday: TrainingWeekday,
    extra_session_dates: Iterable[date] = (),
) -> Season:
    """Return the season for ``weekday``: its sessions numbered ``1..N``, and the total the
    sheet declares for itself.

    ``extra_session_dates`` are merged in on top of what the sheet colours, for days the
    spreadsheet is known to have mis-styled. They are validated like any other date.
    """
    sheet = _open_schedule_sheet(schedule_file)
    session_colour = _find_legend_colour(sheet, weekday)
    blocks = _find_month_blocks(sheet)
    if not blocks:
        raise ScheduleReadError("no month headings found — the sheet layout is not recognised")
    coloured_dates = _collect_coloured_dates(sheet, blocks, session_colour, weekday)
    covered_months = {(block.year, block.month) for block in blocks}
    season_dates = _merge_extra_dates(coloured_dates, extra_session_dates, weekday, covered_months)
    sessions = number_sessions(season_dates)
    declared_total = _read_declared_total(sheet, weekday)
    _verify_against_declared_total(declared_total, weekday, len(sessions))
    return Season(sessions=sessions, declared_total=declared_total)


def _open_schedule_sheet(schedule_file: Path) -> Worksheet:
    try:
        workbook = load_workbook(filename=schedule_file, data_only=True)
    except OSError as error:
        raise ScheduleReadError(f"cannot open schedule file {schedule_file}") from error
    sheet = workbook.active
    if sheet is None:
        raise ScheduleReadError(f"schedule file {schedule_file} has no active sheet")
    return sheet


def _find_legend_colour(sheet: Worksheet, weekday: TrainingWeekday) -> str:
    """Read the colour the legend assigns to ``weekday``, rather than assuming one."""
    matches = [
        colour
        for cell in _all_cells(sheet)
        if _text_of(cell) == weekday.legend_label and (colour := _fill_colour(cell)) is not None
    ]
    unique = set(matches)
    if not unique:
        raise ScheduleReadError(
            f"no filled legend cell labelled {weekday.legend_label!r} — "
            "cannot tell which colour marks this weekday"
        )
    if len(unique) > 1:
        listed = ", ".join(sorted(unique))
        raise ScheduleValidationError(
            f"legend label {weekday.legend_label!r} appears in conflicting colours: {listed}"
        )
    return unique.pop()


def _collect_coloured_dates(
    sheet: Worksheet,
    blocks: list[MonthBlock],
    session_colour: str,
    weekday: TrainingWeekday,
) -> list[date]:
    return [
        day for block in blocks for day in _read_block_dates(sheet, block, session_colour, weekday)
    ]


@dataclass(frozen=True, slots=True)
class MonthHeading:
    """A ``SEPTIEMBRE 2026``-style heading and where it sits."""

    month: int
    year: int
    row: int
    column: int


def _find_month_blocks(sheet: Worksheet) -> list[MonthBlock]:
    """Locate every month grid by its heading, and the band of day rows beneath it."""
    headings = _find_month_headings(sheet)
    heading_rows = sorted({heading.row for heading in headings})
    band_end = {
        row: heading_rows[index + 1] - 1 if index + 1 < len(heading_rows) else sheet.max_row
        for index, row in enumerate(heading_rows)
    }
    return [
        MonthBlock(
            month=heading.month,
            year=heading.year,
            heading_row=heading.row,
            first_column=heading.column,
            first_day_row=heading.row + 2,
            last_day_row=min(band_end[heading.row], heading.row + 1 + MAX_WEEK_ROWS),
        )
        for heading in headings
    ]


def _find_month_headings(sheet: Worksheet) -> list[MonthHeading]:
    headings: list[MonthHeading] = []
    for cell in _all_cells(sheet):
        match = _MONTH_HEADER.match(_text_of(cell))
        if match is None or cell.row is None or cell.column is None:
            continue
        headings.append(
            MonthHeading(
                month=SPANISH_MONTHS[match["month"].upper()],
                year=int(match["year"]),
                row=cell.row,
                column=cell.column,
            )
        )
    return headings


def _read_block_dates(
    sheet: Worksheet, block: MonthBlock, session_colour: str, weekday: TrainingWeekday
) -> list[date]:
    columns_by_weekday = _map_weekday_columns(sheet, block)
    target_column = columns_by_weekday.get(weekday.weekday_index)
    if target_column is None:
        raise ScheduleReadError(
            f"the {block.year}-{block.month:02d} grid has no {weekday.legend_label} column"
        )
    return [
        day
        for row in range(block.first_day_row, block.last_day_row + 1)
        if (day := _read_day_cell(sheet.cell(row=row, column=target_column), block, weekday))
        is not None
        and _fill_colour(sheet.cell(row=row, column=target_column)) == session_colour
    ]


def _map_weekday_columns(sheet: Worksheet, block: MonthBlock) -> dict[int, int]:
    """Map each weekday index to its column, from the ``L M X J V S D`` header row."""
    columns: dict[int, int] = {}
    for column in block.columns:
        initial = _text_of(sheet.cell(row=block.weekday_header_row, column=column))
        weekday_index = SPANISH_WEEKDAY_INITIALS.get(initial.upper()) if initial else None
        if weekday_index is not None:
            columns[weekday_index] = column
    if len(columns) != DAYS_PER_WEEK:
        raise ScheduleReadError(
            f"the {block.year}-{block.month:02d} grid has {len(columns)} weekday columns, "
            f"expected {DAYS_PER_WEEK}"
        )
    return columns


def _read_day_cell(cell: AnyCell, block: MonthBlock, weekday: TrainingWeekday) -> date | None:
    """Turn a day-number cell into a date, refusing anything the grid contradicts."""
    if not isinstance(cell.value, int):
        return None
    try:
        day = date(block.year, block.month, cell.value)
    except ValueError as error:
        raise ScheduleValidationError(
            f"cell {cell.coordinate} holds day {cell.value}, "
            f"which is not a date in {block.year}-{block.month:02d}"
        ) from error
    if not weekday.matches(day):
        raise ScheduleValidationError(
            f"cell {cell.coordinate} sits in the {weekday.legend_label} column but "
            f"{day.isoformat()} is not a {weekday.legend_label} — the grid was misread"
        )
    return day


def _merge_extra_dates(
    coloured_dates: list[date],
    extra_session_dates: Iterable[date],
    weekday: TrainingWeekday,
    covered_months: set[tuple[int, int]],
) -> list[date]:
    """Add hand-supplied dates for days the sheet is known to have mis-styled.

    An override that silently lands on the wrong weekday, outside the season, or twice is
    worse than no override at all, so each is rejected with a reason.
    """
    merged = list(coloured_dates)
    known = set(coloured_dates)
    for day in extra_session_dates:
        if not weekday.matches(day):
            raise ScheduleValidationError(
                f"extra session date {day.isoformat()} is not a {weekday.legend_label}"
            )
        if (day.year, day.month) not in covered_months:
            raise ScheduleValidationError(
                f"extra session date {day.isoformat()} falls outside the months the sheet covers"
            )
        if day in known:
            raise ScheduleValidationError(
                f"extra session date {day.isoformat()} is already a session"
            )
        merged.append(day)
        known.add(day)
    return merged


def _verify_against_declared_total(
    declared: int | None, weekday: TrainingWeekday, found: int
) -> None:
    """Cross-check the count against the total the sheet states in its own subtitle."""
    if declared is not None and declared != found:
        raise ScheduleValidationError(
            f"the sheet declares {declared} {weekday.legend_label.lower()} sessions "
            f"but {found} were found — check EXTRA_SESSION_DATES"
        )


def _read_declared_total(sheet: Worksheet, weekday: TrainingWeekday) -> int | None:
    pattern = re.compile(rf"(\d+)\s+{re.escape(weekday.legend_label.lower())}\b", re.IGNORECASE)
    for cell in _all_cells(sheet):
        text = _text_of(cell)
        if text and (match := pattern.search(text)):
            return int(match[1])
    return None


def _all_cells(sheet: Worksheet) -> Iterable[AnyCell]:
    for row in sheet.iter_rows():
        yield from row


def _text_of(cell: AnyCell) -> str:
    return cell.value.strip() if isinstance(cell.value, str) else ""


def _fill_colour(cell: AnyCell) -> str | None:
    """The cell's solid fill as a six-digit RGB string, or ``None`` if it has no solid fill."""
    fill = cell.fill
    if fill is None or fill.patternType != "solid":
        return None
    rgb = fill.start_color.rgb
    if not isinstance(rgb, str):
        return None
    return rgb[-6:].upper()
